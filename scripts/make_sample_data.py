"""Generate the demo data set: pricing workbook, the client's "existing" quote
sheet, four sample contracts (PDF) and a mock CRM.

Everything here stands in for client material that would arrive in milestone M1.
Run:  python scripts/make_sample_data.py
"""

from __future__ import annotations

import csv
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from closing_agent.pricing import PriceMatrix  # noqa: E402

DATA = ROOT / "data"
CONTRACTS = DATA / "contracts"

PRODUCTS = [
    # sku, name, billing, list_price, highlights
    ("CRM-STD", "Closer CRM Standard", "monthly", 29, "Pipeline, email sync, 5 GB storage"),
    ("CRM-PRO", "Closer CRM Pro", "monthly", 49, "Standard + automation, forecasting, 50 GB"),
    ("CRM-ENT", "Closer CRM Enterprise", "monthly", 89, "Pro + SSO, audit log, dedicated CSM, 99.9% SLA"),
    ("ADD-AI", "AI Assistant add-on", "monthly", 15, "Email drafting, call summaries"),
    ("SVC-ONB", "Onboarding package", "one_off", 1500, "Data migration + 2 training sessions"),
]
VOLUME_TIERS = [
    ("CRM-STD", 1, 0), ("CRM-STD", 10, 5), ("CRM-STD", 25, 10), ("CRM-STD", 50, 15),
    ("CRM-PRO", 1, 0), ("CRM-PRO", 10, 5), ("CRM-PRO", 25, 10), ("CRM-PRO", 50, 15),
    ("CRM-ENT", 1, 0), ("CRM-ENT", 20, 8), ("CRM-ENT", 100, 15),
    ("ADD-AI", 1, 0), ("ADD-AI", 25, 10),
]
TERMS = [(1, 0), (12, 10), (24, 15), (36, 20)]

# The client's existing quote sheet: (ref, sku, qty, term)
REFERENCE_QUOTES = [
    ("Q-2301", "CRM-STD", 8, 12), ("Q-2302", "CRM-STD", 30, 24), ("Q-2303", "CRM-PRO", 12, 12),
    ("Q-2304", "CRM-PRO", 55, 36), ("Q-2305", "CRM-PRO", 3, 1), ("Q-2306", "CRM-ENT", 20, 24),
    ("Q-2307", "CRM-ENT", 120, 36), ("Q-2308", "ADD-AI", 26, 12), ("Q-2309", "ADD-AI", 7, 24),
    ("Q-2310", "SVC-ONB", 1, 12), ("Q-2311", "CRM-STD", 49, 36), ("Q-2312", "CRM-ENT", 7, 12),
]

TEMPLATE_FIELDS = {
    1: [("customer_name", "Customer name"), ("company_number", "Company number"),
        ("billing_email", "Billing email"), ("start_date", "Start date")],
    2: [("sku", "Product SKU"), ("quantity", "Quantity"), ("term_months", "Term (months)"),
        ("total_net", "Total net (GBP)")],
    3: [("sig_customer", "Customer signature"), ("sig_customer_date", "Date signed"),
        ("sig_provider", "Provider signature")],
}


def make_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.append(["sku", "name", "billing", "list_price", "highlights"])
    for row in PRODUCTS:
        ws.append(list(row))
    ws = wb.create_sheet("VolumeTiers")
    ws.append(["sku", "min_qty", "discount_pct"])
    for row in VOLUME_TIERS:
        ws.append(list(row))
    ws = wb.create_sheet("TermDiscounts")
    ws.append(["term_months", "discount_pct"])
    for row in TERMS:
        ws.append(list(row))
    wb.save(path)


def spreadsheet_total(sku: str, qty: int, term: int) -> float:
    """How the client's Excel sheet computes it: plain floats, rounded once at the end.

    Deliberately independent of the engine (which rounds the unit price to pence
    first) so the variance report compares two genuinely different calculations.
    """
    product = next(p for p in PRODUCTS if p[0] == sku)
    vol = max((pct for s, q, pct in VOLUME_TIERS if s == sku and qty >= q), default=0)
    if product[2] == "one_off":
        return round(product[3] * (1 - vol / 100) * qty, 2)
    term_pct = dict(TERMS)[term]
    return round(product[3] * (1 - vol / 100) * (1 - term_pct / 100) * qty * term, 2)


def make_reference_csv(path: Path) -> None:
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["ref", "sku", "quantity", "term_months", "total_net"])
        for ref, sku, qty, term in REFERENCE_QUOTES:
            w.writerow([ref, sku, qty, term, f"{spreadsheet_total(sku, qty, term):.2f}"])


def _header(c: canvas.Canvas, page: int, title: str) -> None:
    c.setFont("Helvetica-Bold", 14)
    c.drawString(50, 800, "Master Subscription Agreement v3")
    c.setFont("Helvetica", 11)
    c.drawString(50, 780, f"Page {page} of 3 - {title}")


def make_form_contract(path: Path, values: dict[str, str]) -> None:
    """Fillable PDF (AcroForm), as produced by e-sign tools or Word/Acrobat forms."""
    c = canvas.Canvas(str(path), pagesize=A4)
    titles = {1: "Customer details", 2: "Order", 3: "Signatures"}
    for page in (1, 2, 3):
        _header(c, page, titles[page])
        y = 730
        for name, label in TEMPLATE_FIELDS[page]:
            c.drawString(50, y + 5, f"{label}:")
            c.acroForm.textfield(name=name, value=values.get(name, ""), x=220, y=y, width=280, height=20,
                                 borderWidth=1, fontSize=10)
            y -= 45
        c.drawString(50, 120, "Initials:")
        c.acroForm.textfield(name=f"init_p{page}", value=values.get(f"init_p{page}", ""), x=220, y=115,
                             width=60, height=20, borderWidth=1, fontSize=10)
        c.showPage()
    c.save()


def make_flat_contract(path: Path, values: dict[str, str]) -> None:
    """Flattened text PDF, as exported from Google Docs: blanks are underscores."""
    c = canvas.Canvas(str(path), pagesize=A4)
    titles = {1: "Customer details", 2: "Order", 3: "Signatures"}
    for page in (1, 2, 3):
        _header(c, page, titles[page])
        y = 730
        for name, label in TEMPLATE_FIELDS[page]:
            c.drawString(50, y, f"{label}: {values.get(name) or '______________________'}")
            y -= 30
        c.drawString(50, 120, f"Initials: {values.get(f'init_p{page}') or '______'}")
        c.showPage()
    c.save()


def main() -> None:
    CONTRACTS.mkdir(parents=True, exist_ok=True)
    xlsx = DATA / "pricing_matrix.xlsx"
    make_workbook(xlsx)
    make_reference_csv(DATA / "reference_quotes.csv")
    matrix = PriceMatrix.from_xlsx(xlsx)

    def full(company: str, email: str, sku: str, qty: int, term: int, total: str | None = None) -> dict:
        total = total or str(matrix.quote(sku, qty, term).total_net)
        return {
            "customer_name": company, "company_number": "09876543", "billing_email": email,
            "start_date": "2026-11-01", "sku": sku, "quantity": str(qty), "term_months": str(term),
            "total_net": total, "sig_customer": "/s/ J. Carter", "sig_customer_date": "2026-10-12",
            "sig_provider": "/s/ M. Reed", "init_p1": "JC", "init_p2": "JC", "init_p3": "JC",
        }

    complete = full("Northbay Logistics Ltd", "accounts@northbay.example", "CRM-PRO", 30, 24)
    make_form_contract(CONTRACTS / "D-1001_northbay.pdf", complete)

    missing = full("Wearside Dental Group", "", "CRM-STD", 12, 12)
    missing.update(sig_customer="", init_p2="", sig_customer_date="")
    make_form_contract(CONTRACTS / "D-1002_wearside.pdf", missing)

    true_total = matrix.quote("CRM-ENT", 25, 36).total_net
    wrong = full("Tyne Freight plc", "finance@tynefreight.example", "CRM-ENT", 25, 36,
                 total=str((true_total * 92 / 100).quantize(true_total)))  # 8% under the matrix
    make_form_contract(CONTRACTS / "D-1003_tynefreight.pdf", wrong)

    gdoc = full("Roker Studios", "hello@roker.example", "CRM-STD", 6, 12)
    gdoc.update(company_number="", init_p3="", sig_provider="")
    make_flat_contract(CONTRACTS / "D-1004_roker_gdoc.pdf", gdoc)

    deals = {
        "D-1001": {"company": "Northbay Logistics Ltd", "contact": "Jo Carter",
                   "email": "jo.carter@northbay.example", "sku": "CRM-PRO", "quantity": 30,
                   "term_months": 24, "stage": "contract_sent", "template": "msa_v3",
                   "contract": "data/contracts/D-1001_northbay.pdf"},
        "D-1002": {"company": "Wearside Dental Group", "contact": "Sam Patel",
                   "email": "sam.patel@wearside.example", "sku": "CRM-STD", "quantity": 12,
                   "term_months": 12, "stage": "contract_sent", "template": "msa_v3",
                   "contract": "data/contracts/D-1002_wearside.pdf"},
        "D-1003": {"company": "Tyne Freight plc", "contact": "Alex Kim",
                   "email": "alex.kim@tynefreight.example", "sku": "CRM-ENT", "quantity": 25,
                   "term_months": 36, "stage": "contract_sent", "template": "msa_v3",
                   "contract": "data/contracts/D-1003_tynefreight.pdf"},
        "D-1004": {"company": "Roker Studios", "contact": "Robin Hall",
                   "email": "robin@roker.example", "sku": "CRM-STD", "quantity": 8,
                   "term_months": 12, "stage": "verbal_yes", "template": "msa_v3",
                   "contract": "data/contracts/D-1004_roker_gdoc.pdf"},
    }
    (DATA / "crm.json").write_text(json.dumps({"deals": deals, "notes": [], "tasks": []}, indent=2),
                                   encoding="utf-8")
    print(f"Sample data written to {DATA}")


if __name__ == "__main__":
    main()
