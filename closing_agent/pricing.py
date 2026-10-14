"""Deterministic pricing engine.

Prices are never computed by the LLM. This module reads the client's pricing
matrix (Excel workbook) and does all arithmetic with Decimal; Claude only
presents the numbers it gets back from here.

Workbook layout (one sheet each, header row first):
  Products       sku | name | billing | list_price | highlights
  VolumeTiers    sku | min_qty | discount_pct
  TermDiscounts  term_months | discount_pct
billing is "monthly" (per unit per month) or "one_off".
"""

from __future__ import annotations

import csv
from dataclasses import asdict, dataclass, field
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

CENT = Decimal("0.01")


def _d(value) -> Decimal:
    return Decimal(str(value))


def _money(value: Decimal) -> Decimal:
    return value.quantize(CENT, rounding=ROUND_HALF_UP)


class PricingError(ValueError):
    """Raised for unknown SKUs, invalid quantities or unsupported terms."""


def parse_money(value) -> Decimal:
    """Parse an amount as people and documents write it: '£15,976.80', ' 15976.8 ', 15976.8."""
    if isinstance(value, Decimal):
        return value
    text = str(value).replace(",", "").replace("£", "").replace("GBP", "").strip()
    try:
        amount = Decimal(text)
    except InvalidOperation:
        raise PricingError(f"{value!r} is not an amount") from None
    if not amount.is_finite():
        raise PricingError(f"{value!r} is not an amount")
    return amount


@dataclass(frozen=True)
class Product:
    sku: str
    name: str
    billing: str  # "monthly" | "one_off"
    list_price: Decimal
    highlights: str = ""


@dataclass
class Quote:
    sku: str
    name: str
    billing: str
    quantity: int
    term_months: int
    list_price: Decimal
    volume_discount_pct: Decimal
    term_discount_pct: Decimal
    extra_discount_pct: Decimal
    unit_net: Decimal
    monthly_net: Decimal | None
    total_net: Decimal
    vat: Decimal
    total_gross: Decimal

    def to_dict(self) -> dict:
        return {k: (str(v) if isinstance(v, Decimal) else v) for k, v in asdict(self).items()}


@dataclass
class PriceMatrix:
    products: dict[str, Product]
    volume_tiers: dict[str, list[tuple[int, Decimal]]]  # sku -> [(min_qty, pct)] ascending
    term_discounts: dict[int, Decimal]
    vat_rate: Decimal = Decimal("0.20")
    currency: str = "GBP"
    source: str = ""
    _catalog_cache: str | None = field(default=None, repr=False)

    # ---------- loading ----------

    @classmethod
    def from_xlsx(cls, path: str | Path, vat_rate: Decimal | float = Decimal("0.20")) -> PriceMatrix:
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            products = {}
            for row in _rows(wb, "Products"):
                sku = str(row["sku"]).strip()
                billing = str(row["billing"]).strip().lower()
                if billing not in ("monthly", "one_off"):
                    raise PricingError(f"{sku}: billing must be 'monthly' or 'one_off', got {billing!r}")
                products[sku] = Product(
                    sku=sku,
                    name=str(row["name"]).strip(),
                    billing=billing,
                    list_price=_d(row["list_price"]),
                    highlights=str(row.get("highlights") or "").strip(),
                )
            tiers: dict[str, list[tuple[int, Decimal]]] = {}
            for row in _rows(wb, "VolumeTiers"):
                sku = str(row["sku"]).strip()
                if sku not in products:
                    raise PricingError(f"VolumeTiers references unknown SKU {sku!r}")
                tiers.setdefault(sku, []).append((int(row["min_qty"]), _d(row["discount_pct"])))
            for sku in tiers:
                tiers[sku].sort()
            terms = {int(r["term_months"]): _d(r["discount_pct"]) for r in _rows(wb, "TermDiscounts")}
        finally:
            wb.close()
        if not terms:
            raise PricingError("TermDiscounts sheet is empty")
        return cls(products, tiers, terms, vat_rate=_d(vat_rate), source=str(path))

    # ---------- core maths ----------

    def product(self, sku: str) -> Product:
        try:
            return self.products[sku]
        except KeyError:
            raise PricingError(f"Unknown SKU {sku!r}. Known: {', '.join(sorted(self.products))}") from None

    def volume_discount(self, sku: str, quantity: int) -> Decimal:
        pct = Decimal("0")
        for min_qty, tier_pct in self.volume_tiers.get(sku, []):
            if quantity >= min_qty:
                pct = tier_pct
        return pct

    def quote(self, sku: str, quantity: int, term_months: int, extra_discount_pct: Decimal | float = 0) -> Quote:
        product = self.product(sku)
        if quantity < 1:
            raise PricingError("quantity must be at least 1")
        extra = _d(extra_discount_pct)
        if not Decimal("0") <= extra < Decimal("100"):
            raise PricingError("extra_discount_pct must be between 0 and 100")

        vol = self.volume_discount(sku, quantity)
        if product.billing == "monthly":
            if term_months not in self.term_discounts:
                raise PricingError(
                    f"Unsupported term {term_months}. Allowed: {sorted(self.term_discounts)}"
                )
            term_pct = self.term_discounts[term_months]
        else:
            term_pct = Decimal("0")  # one-off items are not discounted by contract term

        factor = (1 - vol / 100) * (1 - term_pct / 100) * (1 - extra / 100)
        unit_net = _money(product.list_price * factor)
        if product.billing == "monthly":
            monthly = _money(unit_net * quantity)
            total = _money(monthly * term_months)
        else:
            monthly = None
            total = _money(unit_net * quantity)
        vat = _money(total * self.vat_rate)
        return Quote(
            sku=sku,
            name=product.name,
            billing=product.billing,
            quantity=quantity,
            term_months=term_months,
            list_price=product.list_price,
            volume_discount_pct=vol,
            term_discount_pct=term_pct,
            extra_discount_pct=extra,
            unit_net=unit_net,
            monthly_net=monthly,
            total_net=total,
            vat=vat,
            total_gross=total + vat,
        )

    # ---------- sales helpers ----------

    def price_options(self, sku: str, quantity: int) -> dict:
        """All term options for a line, plus a volume upsell when the next tier is close."""
        product = self.product(sku)
        terms = sorted(self.term_discounts) if product.billing == "monthly" else [0]
        options = [self.quote(sku, quantity, t).to_dict() for t in terms]
        result: dict = {"sku": sku, "name": product.name, "quantity": quantity, "options": options}

        if product.billing == "monthly":
            cheapest = min(options, key=lambda o: Decimal(o["unit_net"]))
            result["lowest_unit_price_term"] = cheapest["term_months"]
            nxt = next((q for q, _ in self.volume_tiers.get(sku, []) if q > quantity), None)
            if nxt is not None and nxt - quantity <= max(2, quantity // 5):
                best_term = cheapest["term_months"]
                current = self.quote(sku, quantity, best_term)
                upsell = self.quote(sku, nxt, best_term)
                result["volume_upsell"] = {
                    "quantity": nxt,
                    "term_months": best_term,
                    "unit_net": str(upsell.unit_net),
                    "unit_saving": str(current.unit_net - upsell.unit_net),
                    "extra_monthly_cost": str(upsell.monthly_net - current.monthly_net),
                }
        return result

    def compare(self, sku: str, quantity: int, term_months: int, offered_total_net, tolerance_pct=1.0) -> dict:
        """Compare a price found elsewhere (contract, email, old sheet) against the matrix."""
        expected = self.quote(sku, quantity, term_months).total_net
        offered = parse_money(offered_total_net)
        variance_pct = (offered - expected) / expected * 100 if expected else Decimal("0")
        ok = abs(variance_pct) <= _d(tolerance_pct)
        return {
            "expected_total_net": str(expected),
            "offered_total_net": str(_money(offered)),
            "difference": str(_money(offered - expected)),
            "variance_pct": str(variance_pct.quantize(Decimal("0.001"))),
            "within_tolerance": ok,
            "tolerance_pct": str(_d(tolerance_pct)),
        }

    def catalog_text(self) -> str:
        """Stable plain-text catalog for the system prompt (deterministic order => cacheable)."""
        if self._catalog_cache is None:
            lines = [f"Currency: {self.currency}. VAT {self.vat_rate * 100:.0f}% added on top of net prices."]
            for sku in sorted(self.products):
                p = self.products[sku]
                unit = "per unit per month" if p.billing == "monthly" else "one-off per unit"
                tiers = ", ".join(f"{q}+ units: -{pct}%" for q, pct in self.volume_tiers.get(sku, []) if pct)
                lines.append(f"- {sku} | {p.name} | list {p.list_price} {unit}"
                             + (f" | volume: {tiers}" if tiers else "")
                             + (f" | {p.highlights}" if p.highlights else ""))
            terms = ", ".join(f"{t} months: -{pct}%" for t, pct in sorted(self.term_discounts.items()))
            lines.append(f"Contract terms (monthly products only): {terms}")
            self._catalog_cache = "\n".join(lines)
        return self._catalog_cache


def _rows(wb, sheet: str):
    if sheet not in wb.sheetnames:
        raise PricingError(f"Workbook is missing sheet {sheet!r}")
    rows = wb[sheet].iter_rows(values_only=True)
    header = [str(h).strip() for h in next(rows)]
    for values in rows:
        if values is None or all(v is None for v in values):
            continue
        yield dict(zip(header, values, strict=False))


# ---------- acceptance criterion 2: variance against the client's own spreadsheet ----------

@dataclass
class VarianceRow:
    ref: str
    sku: str
    quantity: int
    term_months: int
    spreadsheet_total: Decimal
    engine_total: Decimal
    variance_pct: Decimal
    ok: bool


def validate_against_spreadsheet(matrix: PriceMatrix, csv_path: str | Path, tolerance_pct=1.0) -> list[VarianceRow]:
    """Replay every row of the client's existing quote sheet through the engine.

    CSV columns: ref, sku, quantity, term_months, total_net
    """
    rows = []
    with open(csv_path, newline="", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            q = matrix.quote(r["sku"], int(r["quantity"]), int(r["term_months"]))
            sheet_total = _d(r["total_net"])
            var = (q.total_net - sheet_total) / sheet_total * 100 if sheet_total else Decimal("0")
            rows.append(VarianceRow(
                ref=r["ref"], sku=r["sku"], quantity=int(r["quantity"]), term_months=int(r["term_months"]),
                spreadsheet_total=sheet_total, engine_total=q.total_net,
                variance_pct=var.quantize(Decimal("0.0001")), ok=abs(var) <= _d(tolerance_pct),
            ))
    return rows
